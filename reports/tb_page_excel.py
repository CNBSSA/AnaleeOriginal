"""Trial balance Excel export matching the on-screen report table.

Uses ``TrialBalanceContext`` from ``load_trial_balance`` — same accounts, signed
amounts and totals as ``/trial-balance``. BooksXperts intake uses the separate
``build_booksxperts_trial_balance_xlsx`` helper (signed Amount column).
"""
from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from reports.trial_balance_service import TrialBalanceContext

PAGE_TB_COLUMNS = ('Account Code', 'Account Name', 'Debit', 'Credit')


def _debit_credit(amount: Decimal) -> tuple[float, float]:
    if amount > 0:
        return float(amount), 0.0
    if amount < 0:
        return 0.0, float(-amount)
    return 0.0, 0.0


def build_trial_balance_page_xlsx(
    ctx: TrialBalanceContext,
    *,
    company_name: str = '',
    period_end_inclusive: datetime | None = None,
) -> bytes:
    """Workbook with the same rows and totals as the trial balance HTML table."""
    amounts_by_link = {row.link: row.amount for row in ctx.rows}
    wb = Workbook()
    ws = wb.active
    ws.title = 'Trial Balance'

    ws.append(list(PAGE_TB_COLUMNS))
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for account in ctx.accounts:
        balance = amounts_by_link.get(account.link, Decimal('0'))
        debit, credit = _debit_credit(balance)
        ws.append([account.link, account.name, debit, credit])

    total_row_idx = ws.max_row + 2
    ws.cell(row=total_row_idx, column=1, value='Total')
    ws.cell(row=total_row_idx, column=3, value=float(ctx.total_debits))
    ws.cell(row=total_row_idx, column=4, value=float(ctx.total_credits))
    for col in (1, 3, 4):
        ws.cell(row=total_row_idx, column=col).font = Font(bold=True)

    if company_name or period_end_inclusive:
        meta_row = total_row_idx + 2
        meta = 'Analee trial balance'
        if company_name:
            meta += f' — {company_name}'
        if period_end_inclusive:
            meta += f' — {period_end_inclusive.strftime("%Y-%m-%d")}'
        ws.cell(row=meta_row, column=1, value=meta)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
